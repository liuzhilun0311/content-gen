"""Export an illustrated Markdown article without changing its content."""

import argparse
import base64
import hashlib
import html
import json
import io
import mimetypes
import re
import shutil
import subprocess
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from PIL import Image
from task_paths import LONGFORM_VARIANTS, PLATFORMS, validate_task_name
from document_layout import normalize_layout
from markdown_contract import measure, parse


def plain(tokens):
    return "".join(
        "" if t["type"] in ("image", "html") else "\n" if t["type"] == "br"
        else plain(t["tokens"]) if t.get("tokens") else t.get("text", "")
        for t in tokens
    )


def preflight(ast):
    for block in ast:
        kind = block["type"]
        if kind not in ("space", "heading", "paragraph", "text", "blockquote", "list", "code", "table", "hr"):
            raise ValueError("Unsupported article structure before export: " + kind)
        if kind == "code":
            continue
        for token in block.get("tokens", []):
            if kind == "blockquote":
                preflight(block["tokens"])
                break
            check_inline(token)
        if kind == "list":
            for item in block["items"]:
                preflight(item["tokens"])
        if kind == "table":
            for row in [block["header"]] + block["rows"]:
                for cell in row:
                    for token in cell["tokens"]:
                        check_inline(token)


def check_inline(token):
    kind = token["type"]
    if kind == "html":
        if not re.fullmatch(r"</?(?:sup|sub)\s*>|<br\s*/?>", token["text"], re.I):
            raise ValueError("Only inline sup/sub/br HTML is supported; remove layout HTML before export")
    elif kind in ("text", "escape"):
        text = token.get("text", "")
        if re.search(r"\$\$|\\[\[(]|(?<!\\)\$[^$\n]+\$", text):
            raise ValueError("LaTeX math requires verified rendering or equivalent readable Unicode/sup/sub before export")
    elif kind not in ("strong", "em", "image", "link", "br", "codespan"):
        raise ValueError("Unsupported inline article structure: " + kind)
    for child in token.get("tokens", []):
        check_inline(child)


def export(args):
    source = Path(args.source).resolve(strict=True)
    validate_task_name(args.task_name)
    ast = parse(source, args.node, args.marked)
    preflight(ast)
    ast, layout_changes = normalize_layout(ast, args.task_name, getattr(args, "number_sections", False))
    variant = getattr(args, "variant", None)
    text_metrics = measure(ast, variant)
    task_root = getattr(args, "task_root", None)
    if variant and variant not in LONGFORM_VARIANTS:
        raise ValueError("Unknown long-form variant")
    if variant and not task_root:
        raise ValueError("Versioned exports require explicit --task-root")
    is_wechat = args.platform == "wechat"
    is_html = args.platform in ("wechat", "csdn")
    if is_html and not args.summary:
        raise ValueError("HTML export requires a summary based on the approved article")
    if is_wechat and (not args.summary or not args.tags):
        raise ValueError("WeChat requires an article summary and tags based on the approved article")
    destination = Path(args.output_dir).resolve() if args.output_dir else source.parent
    if task_root:
        expected_dir = Path(task_root).resolve() / "publish" / PLATFORMS[args.platform][0]
        if variant:
            expected_dir /= variant
        if destination != expected_dir:
            raise ValueError("Output directory must match task/platform/variant")
    destination.mkdir(parents=True, exist_ok=True)
    platform_label = PLATFORMS[args.platform][0]
    stem = args.task_name + "_" + platform_label + "_publish"
    extension = ".html" if is_html else ".docx"
    output = destination / (stem + extension)
    if output.exists() and not args.replace:
        raise FileExistsError("Fixed output already exists; create a new timestamp task or authorize replacement")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.2)
    for name, size in (("Normal", 12 if is_wechat else 11),
                       ("Title", 19), ("Heading 1", 15), ("Heading 2", 13),
                       ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        pf = style.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(0 if name == "Normal" else 16)
        pf.space_after = Pt(14 if is_wechat else 8)
        pf.line_spacing = 1.75 if is_wechat else 1.25
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.core_properties.author = doc.core_properties.last_modified_by = ""
    image_paths, expected_text, body_html, image_conversions, docx_streams = [], [], [], [], []
    title = ""

    def image_path(token):
        href = token["href"]
        if "://" in href or href.startswith("data:"):
            raise ValueError("Only verified local images can be embedded: " + href)
        path = (source.parent / href).resolve(strict=True)
        target = destination / path.name
        if target.exists():
            if target.read_bytes() != path.read_bytes():
                raise ValueError("Image filename collision: " + str(target))
        elif target != path:
            shutil.copy2(path, target)
        return path

    def inline(paragraph, tokens, bold=False, italic=False, script_mode=None):
        result = []
        for token in tokens:
            kind = token["type"]
            if kind in ("strong", "em"):
                content = inline(paragraph, token["tokens"], bold or kind == "strong",
                                 italic or kind == "em", script_mode)
                tag = "strong" if kind == "strong" else "em"
                result.append(f"<{tag}>{content}</{tag}>")
            elif kind == "image":
                path = image_path(token)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not is_html:
                    docx_stream = None
                    if path.suffix.lower() == ".webp":
                        with Image.open(path) as source_image:
                            original = source_image.convert("RGBA")
                            converted = io.BytesIO()
                            original.save(converted, format="PNG")
                            png_bytes = converted.getvalue()
                        with Image.open(io.BytesIO(png_bytes)) as converted_image:
                            if converted_image.size != original.size or converted_image.convert(
                                    "RGBA").tobytes() != original.tobytes():
                                raise ValueError("WebP to PNG conversion changed pixels: " + str(path))
                        docx_stream = io.BytesIO(png_bytes)
                        docx_streams.append(docx_stream)
                        image_conversions.append({
                            "source": str(path), "embedded_format": "png",
                            "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                            "embedded_sha256": hashlib.sha256(png_bytes).hexdigest(),
                            "dimensions": list(original.size), "rgba_pixels_verified": True,
                        })
                    available_width = section.page_width - section.left_margin - section.right_margin
                    with Image.open(docx_stream if docx_stream else path) as bitmap:
                        ratio = bitmap.width / bitmap.height
                    available_height = section.page_height - section.top_margin - section.bottom_margin
                    width = min(int(available_width * 0.85), int(available_height * 0.70 * ratio))
                    paragraph.add_run().add_picture(
                        docx_stream if docx_stream else str(path), width=width)
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph._p.xpath(".//wp:docPr")[-1].set("descr", token.get("text", ""))
                image_paths.append(path)
                mime = {".png": "image/png", ".jpg": "image/jpeg",
                        ".jpeg": "image/jpeg", ".webp": "image/webp"}.get(
                            path.suffix.lower()) or mimetypes.guess_type(path.name)[0]
                if not mime or not mime.startswith("image/"):
                    raise ValueError("Unsupported image media type: " + path.name)
                uri = "data:" + mime + ";base64," + base64.b64encode(path.read_bytes()).decode("ascii")
                result.append(
                    f'<img src="{uri}" alt="{html.escape(token.get("text", ""), quote=True)}" '
                    'style="display:block;width:85%;max-width:100%;height:auto;margin:20px auto 8px;" />'
                )
            elif kind == "link":
                href = token["href"]
                if not href.startswith(("https://", "http://", "mailto:", "#")):
                    raise ValueError("Unsupported hyperlink: " + href)
                text = plain(token["tokens"])
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(
                    href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                    is_external=True,
                ))
                run, node = OxmlElement("w:r"), OxmlElement("w:t")
                node.text = text
                run.append(node)
                link.append(run)
                paragraph._p.append(link)
                expected_text.append(text)
                result.append(f'<a href="{html.escape(href, quote=True)}" '
                              f'style="color:#333333;text-decoration:underline;">{html.escape(text)}</a>')
            elif kind == "br":
                paragraph.add_run().add_break()
                result.append("<br />")
            elif kind == "html":
                text = token["text"].lower().strip()
                if text.startswith("<br"):
                    paragraph.add_run().add_break()
                    result.append("<br />")
                else:
                    script_mode = None if text.startswith("</") else ("sup" if text.startswith("<sup") else "sub")
                    result.append(text)
            elif kind in ("text", "escape", "codespan"):
                if token.get("tokens"):
                    result.append(inline(paragraph, token["tokens"], bold, italic, script_mode))
                else:
                    text = token.get("text", "")
                    run = paragraph.add_run(text)
                    run.bold, run.italic = bold, italic
                    run.font.superscript = script_mode == "sup"
                    run.font.subscript = script_mode == "sub"
                    expected_text.append(text)
                    result.append(html.escape(text))
            else:
                raise ValueError("Unsupported inline token: " + kind)
        return "".join(result)

    def blocks(items):
        nonlocal title
        result = []
        previous_image_paragraph = None
        for block in items:
            kind = block["type"]
            if kind == "space":
                continue
            if kind == "heading":
                depth = block["depth"]
                paragraph = doc.add_paragraph(style="Title" if depth == 1
                                              else f"Heading {min(depth - 1, 3)}")
                content = inline(paragraph, block["tokens"])
                if depth == 1 and not title:
                    title = plain(block["tokens"])
                else:
                    result.append(f'<h{min(depth, 6)} style="font-family:Microsoft YaHei,'
                                  f'Arial,sans-serif;font-size:20px;line-height:1.5;'
                                  f'font-weight:700;color:#000000;margin:36px 0 16px;'
                                  f'letter-spacing:0;">{content}</h{min(depth, 6)}>')
            elif kind in ("paragraph", "text"):
                p = doc.add_paragraph()
                tokens = block.get("tokens") or [{"type": "text", "text": block.get("text", "")}]
                caption = previous_image_paragraph is not None and (
                    re.match(r"^图[1-9]\d*\s+\S", plain(tokens).strip()) or
                    (len(tokens) == 1 and tokens[0]["type"] == "em"))
                if caption:
                    content = inline(p, [{"type": "text", "text": plain(tokens)}])
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(12)
                    previous_image_paragraph.paragraph_format.keep_with_next = True
                    p.paragraph_format.keep_together = True
                    p.paragraph_format.keep_with_next = False
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        run.italic = False
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(102, 102, 102)
                    result.append('<p style="font-family:Microsoft YaHei,Arial,sans-serif;'
                                  'font-size:14px;font-style:normal;line-height:1.6;color:#666666;'
                                  'text-align:center;text-indent:0;letter-spacing:0;'
                                  'margin:8px 0 24px;overflow-wrap:anywhere;">' + content + "</p>")
                else:
                    content = inline(p, tokens)
                    result.append('<p style="font-family:Microsoft YaHei,Arial,sans-serif;'
                                  'font-size:16px;line-height:1.75;color:#333333;text-align:left;'
                                  'text-indent:0;letter-spacing:0;margin:0 0 20px;">' + content + "</p>")
                previous_image_paragraph = p if any(t["type"] == "image" for t in tokens) else None
            elif kind == "blockquote":
                result.append("<blockquote style=\"margin:16px 0;padding:0;\">" +
                              blocks(block["tokens"]) + "</blockquote>")
            elif kind == "list":
                for index, item in enumerate(block["items"]):
                    label = str(block.get("start", 1) + index) + "." if block["ordered"] else "•"
                    p = doc.add_paragraph()
                    # List numbering is formatting rather than source prose.
                    p.add_run(label + " ")
                    expected_text.append(label + " ")
                    content = []
                    nested = []
                    for child in item["tokens"]:
                        if child["type"] in ("text", "paragraph"):
                            content.append(inline(p, child["tokens"]))
                        else:
                            nested.append(child)
                    result.append(f'<p style="font-size:16px;line-height:1.75;margin:0 0 16px;'
                                  f'color:#333333;">{label} ' + "<br />".join(content) + "</p>")
                    if nested:
                        result.append('<div style="margin-left:20px;">' + blocks(nested) + "</div>")
            elif kind == "code":
                p = doc.add_paragraph()
                p.add_run(block["text"])
                expected_text.append(block["text"])
                result.append('<pre style="white-space:pre-wrap;word-break:break-word;'
                              'font-size:14px;line-height:1.6;">' + html.escape(block["text"]) + "</pre>")
            elif kind == "table":
                rows = [block["header"]] + block["rows"]
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Table Grid"
                result.append('<table style="width:100%;border-collapse:collapse;font-size:14px;">')
                for row_index, cells in enumerate(rows):
                    row = table.add_row()
                    result.append("<tr>")
                    for cell, value in zip(row.cells, cells):
                        content = inline(cell.paragraphs[0], value["tokens"], bold=row_index == 0)
                        result.append('<td style="border:1px solid #d9d9d9;padding:8px;'
                                      'line-height:1.6;word-break:break-word;">' + content + "</td>")
                    result.append("</tr>")
                result.append("</table>")
            elif kind == "hr":
                continue
            else:
                raise ValueError("Unsupported block token: " + kind)
            if kind not in ("paragraph", "text"):
                previous_image_paragraph = None
        return "".join(result)

    body_html.append(blocks(ast))
    media = []
    actual_text = "".join(expected_text)
    if not is_html:
        doc.save(output)
        verify_docx(output, image_paths, expected_text, image_conversions)
        with ZipFile(output) as archive:
            media = [n for n in archive.namelist() if n.startswith("word/media/")]
    html_output = None
    if is_html:
        html_output = output
        wrapper = Path(__file__).resolve().parent.parent / "assets" / "wechat-copy.html"
        page = wrapper.read_text(encoding="utf-8")
        values = {
            "TITLE": html.escape(title), "SUMMARY": html.escape(args.summary),
            "SUMMARY_LABEL": "简介" if is_wechat else "文章摘要",
            "SUMMARY_HEADING": "文章简介" if is_wechat else "文章摘要",
            "BODY": "".join(body_html),
            "TAGS_BUTTON": '<button id="copy-tags" type="button">复制标签</button>' if is_wechat else "",
            "TAGS_BLOCK": (
                '<p style="font-size:13px;color:#666;margin:0 0 8px;">文章标签</p>'
                '<p id="tags" style="font-size:15px;line-height:1.7;margin:0;">'
                + html.escape(args.tags) + "</p>"
            ) if is_wechat else "",
            "TAGS_LISTENER": (
                'document.getElementById("copy-tags").addEventListener("click", event => '
                'copyContent("tags", false, event.currentTarget));'
            ) if is_wechat else "",
        }
        page = re.sub(r"\{\{([A-Z_]+)\}\}",
                      lambda match: values[match.group(1)], page)
        from lxml import html as html_parser
        tree = html_parser.fromstring(page)
        article = tree.get_element_by_id("article")
        exported_text = title + "".join(article.itertext())
        assert "".join(exported_text.split()) == "".join(actual_text.split()), "HTML text differs"
        assert not article.xpath(".//*[@id='summary' or @id='tags' or @id='title']")
        images = article.xpath(".//img")
        assert len(images) == len(image_paths)
        for image, original in zip(images, image_paths):
            assert base64.b64decode(image.get("src").split(",", 1)[1]) == original.read_bytes()
        html_output.write_text(page, encoding="utf-8")
    report = {
        "variant": variant,
        "platform": args.platform, "docx": None if is_html else str(output),
        "html": str(html_output) if html_output else None,
        "summary": args.summary if is_html else None,
        "tags": args.tags if is_wechat else None,
        "image_occurrences": len(image_paths), "embedded_media": len(image_paths) if is_html else len(media),
        "text_characters": len(actual_text), "text_verified": True,
        "images": [str(p) for p in image_paths],
        "image_conversions": image_conversions,
        "target_editor_verified": False,
        "layout_changes": layout_changes,
        "number_sections": getattr(args, "number_sections", False),
        "text_metrics": text_metrics,
        "length_deviation_reason": getattr(args, "length_deviation_reason", None),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "output_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
    }
    root = Path(task_root).resolve() if task_root else destination.parent.parent
    report_path = root / "_work" / "export" / args.platform if variant else root / "_work" / "export"
    report_path.mkdir(parents=True, exist_ok=True)
    (report_path / ((variant if variant else args.platform) + ".json")).write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False))


def verify_docx(output, image_paths, expected_text, image_conversions=()):
    with ZipFile(output) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        hashes = {hashlib.sha256(archive.read(n)).hexdigest() for n in media}
        expected_hashes = {hashlib.sha256(p.read_bytes()).hexdigest() for p in image_paths
                           if p.suffix.lower() != ".webp"}
        expected_hashes.update(item["embedded_sha256"] for item in image_conversions)
        assert expected_hashes <= hashes
        rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        assert all(r.get("TargetMode") != "External" for r in rels
                   if r.get("Type", "").endswith("/image"))
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        actual_text = "".join(xml.xpath("//w:t/text()", namespaces=ns))
        assert actual_text == "".join(expected_text), "Exported text differs"
        blips = xml.xpath("//a:blip", namespaces=ns)
        assert len(blips) == len(image_paths)
        targets = {r.get("Id"): r.get("Target") for r in rels
                   if r.get("Type", "").endswith("/image")}
        embed = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        for blip, original in zip(blips, image_paths):
            target = targets[blip.get(embed)]
            media_path = target.lstrip("/") if target.startswith("/") else "word/" + target
            embedded = archive.read(media_path)
            if original.suffix.lower() == ".webp":
                with Image.open(io.BytesIO(original.read_bytes())) as source_image:
                    expected_pixels, expected_size = source_image.convert("RGBA").tobytes(), source_image.size
                with Image.open(io.BytesIO(embedded)) as embedded_image:
                    assert embedded_image.size == expected_size
                    assert embedded_image.convert("RGBA").tobytes() == expected_pixels, \
                        "Lossless-converted image pixels differ"
            else:
                assert embedded == original.read_bytes(), "Image order or bytes differ"


if __name__ == "__main__":
    from markdown_contract import runtime_paths
    default_node, default_marked = runtime_paths()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", required=True)
    parser.add_argument("--task-name", required=True)
    parser.add_argument("--platform", choices=["zhihu", "wechat", "csdn"], required=True)
    parser.add_argument("--node", default=default_node)
    parser.add_argument("--marked", default=default_marked)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--summary")
    parser.add_argument("--tags")
    parser.add_argument("--replace", action="store_true",
                        help="Replace an existing generated export only when explicitly requested")
    parser.add_argument("--variant", choices=LONGFORM_VARIANTS, help="Historical versioned tasks only")
    parser.add_argument("--task-root", help="Explicit task root for paths and validation records")
    parser.add_argument("--length-deviation-reason", help="Historical compatibility metadata; no length gate")
    export(parser.parse_args())
